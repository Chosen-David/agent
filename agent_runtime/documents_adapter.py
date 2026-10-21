"""Bounded document operations for the opt-in documents environment.

Executed in its managed venv by documents.py. Parser isolation is a resource
boundary, not an OS sandbox or a semantic/visual quality certificate.
"""
import argparse
import hashlib
import importlib.metadata
import json
import math
from pathlib import Path
import tempfile


INPUT_LIMIT = 16 * 1024 * 1024
TEXT_LIMIT = 1024 * 1024
PACKAGES = ('PyMuPDF', 'markdown-it-py', 'python-pptx', 'Jinja2', 'mdurl',
            'lxml', 'MarkupSafe', 'Pillow', 'typing-extensions', 'XlsxWriter')


def bounded_text(value, limit=TEXT_LIMIT):
    if not isinstance(value, str) or len(value.encode('utf-8')) > limit:
        raise ValueError('text exceeds byte limit or has invalid type')
    return value


def latex_escape(text):
    replacements = {'\\': r'\textbackslash{}', '{': r'\{', '}': r'\}',
                    '$': r'\$', '&': r'\&', '#': r'\#', '%': r'\%',
                    '_': r'\_', '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}
    return ''.join(replacements.get(c, c) for c in text)


def execute(operation, payload, output_dir):
    out = Path(output_dir)
    if operation == 'pdf-read':
        import pymupdf
        raw = bytes.fromhex(payload['pdf_hex'])
        if len(raw) > INPUT_LIMIT:
            raise ValueError('PDF exceeds 16 MiB')
        pages = payload.get('pages', [1])
        if (type(pages) is not list or not 1 <= len(pages) <= 16 or
                any(type(n) is not int for n in pages)):
            raise ValueError('pages must contain 1–16 integer page numbers')
        with pymupdf.open(stream=raw, filetype='pdf') as doc:
            if doc.needs_pass or not 1 <= len(doc) <= 1000:
                raise ValueError('encrypted PDF or page count outside 1–1000')
            result, used = [], 0
            for n in pages:
                if not 1 <= n <= len(doc):
                    raise ValueError('page number outside PDF')
                page = doc[n - 1]
                width, height = page.rect.width, page.rect.height
                if not all(math.isfinite(x) and x > 0 for x in (width, height)):
                    raise ValueError('invalid page dimensions')
                scale = min(1.0, 2000 / max(width, height))
                w, h = math.ceil(width * scale), math.ceil(height * scale)
                if w * h > 4_000_000 or max(w, h) > 2000:
                    raise ValueError('render pixel limit')
                text = page.get_text(sort=True)
                used += len(text.encode('utf-8'))
                if used > TEXT_LIMIT:
                    raise ValueError('PDF text exceeds 1 MiB')
                preview = out / f'page-{n}.png'
                page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False,
                                colorspace=pymupdf.csRGB).save(preview)
                result.append(dict(page=n, text=text, preview=preview.name,
                                   preview_sha256=hashlib.sha256(preview.read_bytes()).hexdigest()))
        return dict(pages=result)
    if operation == 'markdown-render':
        from markdown_it import MarkdownIt
        text = bounded_text(payload['text'])
        html = MarkdownIt('commonmark', {'html': False}).render(text)
        bounded_text(html, 4 * TEXT_LIMIT)
        (out / 'document.html').write_text(html, encoding='utf-8')
        return dict(file='document.html', html=html)
    if operation == 'pptx-create':
        from pptx import Presentation
        slides = payload['slides']
        if type(slides) is not list or not 1 <= len(slides) <= 64:
            raise ValueError('slides must contain 1–64 entries')
        presentation = Presentation()
        for item in slides:
            if type(item) is not dict or set(item) != {'title', 'body'}:
                raise ValueError('slide requires exactly title and body')
            title = bounded_text(item['title'], 16000)
            body = bounded_text(item['body'], 16000)
            if len((title + body).encode('utf-8')) > 16000:
                raise ValueError('slide text exceeds 16000 UTF-8 bytes')
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        path = out / 'document.pptx'
        presentation.save(path)
        reopened = Presentation(path)
        if len(reopened.slides) != len(slides):
            raise ValueError('PPTX reopen mismatch')
        return dict(file=path.name, slide_count=len(reopened.slides))
    if operation == 'latex-source':
        from jinja2 import Environment, StrictUndefined
        text = bounded_text(payload['text'], 65536)
        template = Environment(undefined=StrictUndefined, autoescape=False).from_string(
            '\\documentclass{article}\n\\begin{document}\n{{ body }}\n\\end{document}\n')
        source = template.render(body=latex_escape(text))
        (out / 'document.tex').write_text(source, encoding='utf-8')
        return dict(file='document.tex', source=source, latex_compile='unavailable')
    if operation == 'probe':
        import pymupdf
        with tempfile.TemporaryDirectory(dir=out) as temp:
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((72, 72), 'Agent documents smoke')
            raw = document.tobytes()
            document.close()
            pdf = execute('pdf-read', dict(pdf_hex=raw.hex(), pages=[1]), temp)
            if 'Agent documents smoke' not in pdf['pages'][0]['text']:
                raise ValueError('PDF extraction mismatch')
            md = execute('markdown-render', dict(text='# Agent\n\n**smoke**'), temp)
            if '<strong>smoke</strong>' not in md['html']:
                raise ValueError('Markdown conversion mismatch')
            execute('pptx-create', dict(slides=[dict(title='Agent', body='smoke')]), temp)
            tex = execute('latex-source', dict(text='a_b%{x}\\'), temp)
            if r'a\_b\%\{x\}\textbackslash{}' not in tex['source']:
                raise ValueError('LaTeX escaping mismatch')
        return dict(capabilities={k: 'ready' for k in
                    ('pdf_read_render', 'markdown_render', 'pptx_create_reopen', 'latex_source')},
                    latex_compile='unavailable',
                    packages={name: importlib.metadata.version(name) for name in PACKAGES})
    raise ValueError('unsupported operation')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('operation', choices=('probe', 'pdf-read', 'markdown-render',
                                             'pptx-create', 'latex-source'))
    parser.add_argument('input')
    parser.add_argument('output_dir')
    args = parser.parse_args()
    # Applies before importing document parsers. The parent also enforces wall
    # time and a file-backed output cap; no shell or arbitrary generator runs.
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (1024**3, 1024**3))
        resource.setrlimit(resource.RLIMIT_FSIZE, (64 * 1024**2, 64 * 1024**2))
        resource.setrlimit(resource.RLIMIT_CPU, (50, 50))
        raw = Path(args.input).read_bytes()
        if len(raw) > 2 * INPUT_LIMIT + 4096:
            raise ValueError('input envelope exceeds limit')
        payload = json.loads(raw)
        result = dict(schema_version='documents-result/v1', status='ok',
                      operation=args.operation, result=execute(args.operation, payload, args.output_dir))
    except ImportError as exc:
        result = dict(schema_version='documents-result/v1', status='unavailable',
                      operation=args.operation, error=str(exc)[:2000])
    except (ValueError, KeyError, TypeError, OSError, RuntimeError) as exc:
        result = dict(schema_version='documents-result/v1', status='invalid',
                      operation=args.operation, error=str(exc)[:2000])
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result['status'] == 'ok' else (2 if result['status'] == 'unavailable' else 3)


if __name__ == '__main__':
    raise SystemExit(main())
